import io
from pypdf import PdfReader
import docx

class DocumentParserService:
    @staticmethod
    def extract_text_from_pdf(file_bytes: bytes) -> str:
        text = ""
        try:
            reader = PdfReader(io.BytesIO(file_bytes))
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\n"
        except Exception as e:
            text = f"Error reading PDF: {str(e)}"
        return text.strip()

    @staticmethod
    def extract_text_from_docx(file_bytes: bytes) -> str:
        text = ""
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
            for para in doc.paragraphs:
                if para.text:
                    text += para.text + "\n"
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                    if row_text:
                        text += row_text + "\n"
        except Exception as e:
            text = f"Error reading DOCX: {str(e)}"
        return text.strip()

    @staticmethod
    def extract_text_from_txt(file_bytes: bytes) -> str:
        try:
            return file_bytes.decode("utf-8", errors="ignore").strip()
        except Exception as e:
            return f"Error reading TXT: {str(e)}"

    @classmethod
    def parse_file(cls, filename: str, content: bytes) -> str:
        filename_lower = filename.lower()
        if filename_lower.endswith(".pdf"):
            return cls.extract_text_from_pdf(content)
        elif filename_lower.endswith(".docx") or filename_lower.endswith(".doc"):
            return cls.extract_text_from_docx(content)
        elif filename_lower.endswith(".txt"):
            return cls.extract_text_from_txt(content)
        else:
            return content.decode("utf-8", errors="ignore").strip()

parser_service = DocumentParserService()
